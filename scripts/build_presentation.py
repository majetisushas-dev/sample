from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
R=Path(__file__).resolve().parents[1]; out=R/'presentation/zeroth_review.pptx'; prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
slides=[
('Evidence-Calibrated Self-Correcting RAG for Hallucination-Resistant Question Answering','Zeroth review — proposed work\n[Member 1 name] · [Member 2 name] · [Member 3 name]\n[Institution] · [Guide name]'),
('Motivation','A fluent QA system may answer from weak or conflicting passages.\nGrounded scenario: an educational assistant cites a nearby passage that does not support its date claim.\nWe propose to expose evidence, calibrate confidence, and abstain.'),
('Problem statement and scope','Build a local CPU question-answering pipeline for a bounded public corpus.\nDistinguish retrieval relevance, evidence support, and answer correctness.\nExclude paid APIs, publication guarantees, and claims of eliminating hallucination.'),
('Objectives and research questions','RQ1: Will verification plus bounded feedback improve correctness/support?\nRQ2: Will calibration improve ECE and Brier score?\nRQ3: Can abstention lower risk at reported coverage and cost?\nRQ4: Which components contribute most?'),
('Verified literature','RAG: Lewis et al., NeurIPS 2020.\nDPR: Karpukhin et al., EMNLP 2020, DOI 10.18653/v1/2020.emnlp-main.550.\nSelf-RAG: Asai et al., ICLR 2024.\nCRAG: Yan et al., arXiv:2401.15884.\nCalibration: Guo et al., ICML 2017.'),
('Research gap and proposed contribution','Existing components do not automatically yield calibrated, supported answers on a laptop.\nWe propose a transparent bounded policy, logged evidence changes, a locally trained verifier, and selective evaluation.\nContribution is engineering integration and evaluation—not a claimed new learning algorithm.'),
('Proposed system','Question → sparse+dense retrieval → evidence selection → candidate\n→ verifier → one/two retrieval-changing corrections\n→ correctness calibration → cited answer or abstention.'),
('Workflow and stopping','Retrieve and generate deterministically.\nScore claims against cited chunks.\nIf insufficient, reformulate once and retrieve new evidence.\nStop on support, repeated query, or round budget.\nLog rankings, candidates, scores, citations, reasons, and latency.'),
('Training and CPU feasibility','Will freeze retrieval/generation choices before calibration.\nWill train a compact evidence-support classifier and final calibrator on disjoint data.\nTarget: Python 3.11, 8–16 GB RAM, Windows CPU.\nA small extractive fallback will protect schedule.'),
('Dataset, baselines, evaluation','Primary plan: SQuAD 2.0 bounded context corpus with documented license.\nB0 closed book; B1 BM25; B2 hybrid; B3 correction; B4 calibrated abstention.\nEM/F1, Recall@k, citation checks, ECE/Brier, risk–coverage, latency/RAM.\nQuestion-level paired bootstrap; locked test IDs.'),
('Scope, deliverables, exclusions','Planned: source, trained artifact, results folder, local dashboard, six-page paper, and editable PPTX.\nExcluded: paid APIs, mandatory GPU fine-tuning, full-Wikipedia claims, and fabricated human evaluation.'),
('SDG alignment','SDG 4: better evidence visibility may support access to educational information.\nSDG 9: reproducible CPU research infrastructure supports local innovation.\nThese are mechanisms of relevance, not measured societal impacts.'),
('Three-member allocation','Member 1: dataset engineering, retrieval, baselines.\nMember 2: verifier training, calibration, correction policy.\nMember 3: dashboard, experiment reporting, documentation, integration.\nAll: evaluation review, manuscript, reproducibility checks.'),
('Eight-week timeline','W1 scope/licensing → W2 corpus/retrieval → W3 generator/smoke test\n→ W4 verifier → W5 correction/calibration → W6 locked evaluation\n→ W7 dashboard/results/paper → W8 audit, review, contingency.\nDependencies: frozen splits before fitting; frozen pipeline before calibration.'),
('Expected outcomes, risks, success','Expected: runnable CPU prototype, real learned component, transparent abstention curves.\nRisks: weak generation, domain shift, low calibration positives, runtime.\nSuccess: reproducible matched-ID evaluation and honest failure reporting—not a promised accuracy percentage.')]
for i,(title,body) in enumerate(slides):
 s=prs.slides.add_slide(prs.slide_layouts[6]); bg=s.background.fill; bg.solid(); bg.fore_color.rgb=RGBColor(9,20,38)
 tb=s.shapes.add_textbox(Inches(.75),Inches(.55),Inches(11.8),Inches(1)).text_frame; p=tb.paragraphs[0]; p.text=title; p.font.size=Pt(28); p.font.bold=True; p.font.color.rgb=RGBColor(61,214,198)
 bb=s.shapes.add_textbox(Inches(.85),Inches(1.75),Inches(11.5),Inches(4.8)).text_frame; bb.word_wrap=True; p=bb.paragraphs[0]; p.text=body; p.font.size=Pt(22); p.font.color.rgb=RGBColor(235,241,248); p.space_after=Pt(14)
 ft=s.shapes.add_textbox(Inches(.8),Inches(7.05),Inches(11.8),Inches(.25)).text_frame.paragraphs[0]; ft.text=f'Proposal-stage zeroth review · slide {i+1}/15'; ft.font.size=Pt(9); ft.font.color.rgb=RGBColor(145,160,178)
prs.save(out); print(out)

